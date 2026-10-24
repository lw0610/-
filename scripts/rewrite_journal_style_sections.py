from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import RGBColor, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zw/研究生/QGIS-agent/Eino")
SRC = ROOT / "AutoCarto_专题图制图智能体论文_第4-5章优化版.docx"
OUT = ROOT / "AutoCarto_专题图制图智能体论文_期刊风格结构版.docx"

RED = RGBColor(192, 0, 0)


CONTENT = [
    ("3 系统架构与方法", "Heading 1"),
    ("本文提出一种面向路线类专题图的知识约束智能制图方法。该方法针对自然语言路线叙事难以直接转化为 GIS 制图操作，通用大模型缺少路线制图语义约束，自动生成地图缺少质量反馈等问题，将路线语义建模，知识图谱约束，多智能体协同，QGIS 工具执行和质量反馈机制组织为统一框架。与直接让大语言模型生成 PyQGIS 脚本不同，本文方法在工具执行前构造可检查的制图状态，使路线节点，历史事件，空间位置，表达方法，工具输入和评价结果能够在同一任务链中传递与修订。", None),
    ("图 3-1 知识约束的路线类专题图智能制图方法框架", "fig_placeholder"),
    ("3.1 方法总体设计", "Heading 2"),
    ("路线类专题图智能制图任务可以定义为从自然语言路线叙事到可验证地图成果的转换过程。给定用户制图指令，历史文献或业务文本，系统需要识别其中的行动主体，时间线索，地点实体，路线节点，历史事件，节点顺序和成果表达要求，并结合已有空间数据与 QGIS 工程状态，生成包含路线线状符号，节点标注，事件说明，图例，比例尺，指北针和输出布局的专题地图。该任务的输入包括文本需求，空间数据，制图知识，模型配置和工具状态。输出包括结构化制图需求，路线制图知识图谱，工具调用序列，QGIS 图层与布局对象以及最终导出的地图成果。", None),
    ("AutoCarto 智能制图系统采用用户交互层，Go 后端调度层，多智能体层和 QGIS 执行层组成的四层架构。用户交互层负责接收自然语言制图需求，路线文本和制图参数，并展示任务状态，工具日志和地图成果。Go 后端调度层负责维护会话上下文，任务队列，模型配置，工具网关和运行日志，将前端输入转化为可调度的结构化请求。多智能体层负责完成路线意图解析，数据可用性推理，路线关系组织，专题表达设计，版式编排和质量评价。QGIS 执行层以 QGIS 和 PyQGIS 为本地制图环境，将工具调用转化为图层加载，几何处理，符号化，标注配置，布局生成和成果导出等操作。", None),
    ("该总体设计的关键在于区分语言推理，任务调度和 GIS 执行的边界。大语言模型负责理解制图意图，组织路线语义并选择合适工具。Go 后端负责运行时状态维护，工具调用转发和异常记录。QGIS 负责可验证的空间数据处理和地图渲染。通过这种分层设计，系统既避免了模型直接生成长脚本带来的不可控问题，也避免了 GIS 工具孤立执行而缺少语义约束的问题。", None),
    ("3.2 路线制图语义建模与知识组织", "Heading 2"),
    ("路线类专题图的核心信息通常以叙事方式出现，例如某一行动主体在特定时间从起点出发，经过若干地点，并在途中发生战斗，转移，驻留或会师等事件。此类文本不能直接作为 GIS 绘图参数使用，因为地点实体，路线顺序，事件归属和表达意图往往隐含在句子结构中。为此，本文将路线制图信息抽取定义为本体约束下的语义结构化过程，即将原始文本解析为可供知识图谱，智能体规划和工具执行共享的制图需求。", None),
    ("抽取模式包括行动主体，时间信息，地点实体，路线节点，历史事件，路线顺序，节点重要性，事件说明和成果表达要求等字段。地点实体记录名称，别名，地理线索和地理编码需求。路线节点记录节点序号，节点类型，关联地点和标注等级。历史事件记录事件名称，发生时间，关联节点和说明文本。表达要求记录路线线型，节点符号，说明框，图例和导出格式等需求。模型输出需要满足预定义字段集合和类型约束，并经过字段完整性，节点顺序一致性和事件地点匹配检查。", None),
    ("在语义抽取基础上，系统构建面向路线类专题图的制图本体与知识图谱。该知识图谱不是通用地理知识库，而是围绕路线图自动制图任务组织的领域知识结构，用于显式描述制图任务，路线节点，路线段，历史事件，空间位置，行动主体，表达方法，地图要素，质量约束和输出成果之间的关系。通过这些关系，系统能够在工具执行前明确哪些节点需要连接，哪些事件需要标注，哪些符号需要进入图例，以及哪些质量问题需要检查。", None),
    ("3.3 多智能体协同与制图反馈机制", "Heading 2"),
    ("路线类专题图制图链条较长，涉及文本理解，数据检查，路线组织，符号设计，版式编排和质量评价等多类知识。若由单一模型一次性完成，容易出现职责混淆，执行顺序错误和质量问题难以定位等情况。本文依据路线类专题图制图过程设置中央协调智能体，路线意图解析智能体，数据推理智能体，路线表达智能体，版式编排智能体和质量评价智能体，使不同知识类型由相应角色处理。", None),
    ("各智能体之间传递的是结构化制图需求，图谱关系，工具执行结果和评价状态，而不是仅依赖自由文本对话。中央协调智能体维护全局目标，任务状态和子任务顺序。路线意图解析智能体识别制图目标，行动主体，路线范围和成果形式。数据推理智能体检查路线节点，坐标字段，图层状态和 QGIS 工程环境。路线表达智能体依据图谱关系生成主路线，分支路线，多起点汇合路线和事件节点表达方案。版式编排智能体组织标题，图例，比例尺，指北针，节点标注和说明框。质量评价智能体根据知识图谱，QGIS 状态和地图图像输出问题定位与修订建议。", None),
    ("QGIS 制图执行采用工具契约而非自由脚本生成的方式组织。每个工具均限定输入参数，输出结果，适用条件和失败反馈，并由 Go 后端统一调度。工具链包括数据接入与检查，路线关系构建，路线表达，地图要素配置和成果导出等能力。工具执行过程中产生的图层状态，日志，布局对象和导出成果会被写入执行状态，供质量评价模块进一步检查。评价结果随后被转化为修订指令，反馈给路线表达智能体或版式编排智能体，使系统形成生成，评价与修订相结合的闭环制图过程。", None),
    ("4 应用与实验", "Heading 1"),
    ("为检验本文系统在路线类专题图制图任务中的可行性，本文以路线图生成为典型应用场景，围绕自然语言路线叙事，结构化制图知识，QGIS 自动制图和质量评价反馈的完整链路开展实验。实验不以覆盖所有 GIS 制图任务为目标，而是重点考察系统在路线节点抽取，路线关系组织，工具调用执行，地图要素配置和反馈修订等环节中的表现。与只展示最终地图成果不同，本章同时记录中间结构化结果，知识组织结果和执行日志，以说明系统能否将文本中的路线语义稳定传递到地图生成过程。", None),
    ("4.1 实验设置", "Heading 2"),
    ("实验任务以行军路线图制作为代表。该类任务通常包含行动主体，行军时间，途经地点，路线顺序，关键事件和制图表达要求等信息，能够较集中地体现路线类专题图中文本语义，空间关系和地图表达之间的转换问题。实验输入包括路线叙事文本，底图与行政边界数据，节点坐标或可定位地点信息，QGIS 工程环境以及系统模型与工具配置。实验输出包括结构化文本解析结果，路线制图知识图谱，智能体工具调用日志，自动生成地图以及质量评价与修订结果。", None),
    ("为避免仅以个例展示系统功能，实验任务按照路线结构和表达复杂度划分为三类。简单任务对应单一路线和少量关键节点。中等任务包含主路线，分支路线或多个事件说明。复杂任务包含多起点汇合，多段路线关系，事件注记和较完整的版式表达。评价指标分为过程指标和成果指标两类。过程指标包括任务完成率，工具调用成功率，工具调用错误次数，人工干预次数和制图耗时。成果指标包括节点完整性，路线连贯性，事件标注准确性，图例一致性，版式完整性和综合质量评分。实验评价指标及其含义见表 4-1。", None),
    ("表 4-1 路线类专题图自动制图实验评价指标", "table_metrics"),
    ("4.2 实验过程与结果呈现", "Heading 2"),
    ("路线类专题图的制图信息通常并不以标准参数形式给出，而是隐含在历史文献，业务材料或用户自然语言描述中。实验首先对路线文本进行结构化解析，将原始叙事中的行动主体，时间信息，地点实体，路线节点，历史事件，节点顺序和成果表达要求抽取为统一字段。示例文本中，系统需要识别中央红军为行动主体，遵义，桐梓，娄山关和赤水为路线节点，娄山关战斗为历史事件，并形成遵义，桐梓，娄山关和赤水的节点顺序关系。文本中的突出主行军路线，标注关键节点和对娄山关战斗进行说明等描述，则被转换为主路线强调，节点标注和事件说明框配置等制图表达要求。", None),
    ("图 4-1 路线文本输入与结构化解析结果", "fig_placeholder"),
    ("在完成信息抽取后，系统进一步将时间，地点，事件，路线顺序和表达要求组织为路线制图知识图谱。知识图谱显式表达路线任务，路线节点，历史事件，路线关系，表达方法和地图要素之间的关系。通过知识图谱组织，系统将文本中的隐含路线关系显式化，为后续智能体规划和 QGIS 工具调用提供约束。", None),
    ("图 4-2 路线制图知识图谱组织结果", "fig_placeholder"),
    ("在自动制图阶段，中央协调智能体依据制图需求和图谱关系组织任务顺序，并调用相应 QGIS 工具完成数据检查，路线段生成，符号配置，节点标注，事件说明框生成，图例配置，版式组织和成果导出等操作。执行日志用于记录不同智能体的调用顺序，工具输入参数，工具返回状态和异常反馈。与直接生成完整 PyQGIS 脚本相比，基于工具契约的执行方式能够将长流程制图任务拆分为可检查的中间步骤，使错误更容易定位。", None),
    ("图 4-3 智能体协同制图执行日志", "fig_placeholder"),
    ("图 4-4 自动生成的路线类专题图成果", "fig_placeholder"),
    ("质量评价环节将自动生成的初始地图，原始制图需求，路线知识图谱以及 QGIS 图层与布局状态共同作为输入，对路线连贯性，节点完整性，事件标注准确性，符号层次，图例一致性，标注可读性和版式完整性进行检查。评价结果随后被转换为修订指令，反馈给路线表达智能体或版式编排智能体，用于调整线宽，线型，节点标注，说明框位置和图例内容。", None),
    ("图 4-5 路线类专题图质量评价结果", "fig_placeholder"),
    ("图 4-6 反馈修订后的路线类专题图成果", "fig_placeholder"),
    ("当前版本重点保留实验设计，评价指标和结果呈现位置。具体任务数量，对比实验数据，质量评分和修订前后量化结果将在系统进一步稳定后补充。", "red"),
    ("5 讨论", "Heading 1"),
    ("本文系统的核心价值并不在于替代所有 GIS 制图操作，而在于为路线类专题图建立从自然语言到结构化制图过程的中间表示。路线图制作的难点不是单纯绘制线状要素，而是需要同时处理路线节点顺序，历史事件归属，空间位置匹配，专题符号表达和版式组织。若直接由通用大语言模型生成 PyQGIS 脚本，模型需要在一次输出中同时完成需求理解，字段判断，代码生成和版式控制，容易受到图层状态，字段名称，路径配置和参数缺失等因素影响。本文方法通过文本解析和知识图谱组织先形成明确的路线制图知识，再由多智能体选择工具并执行局部操作，从而提高长流程制图任务的可控性。", None),
    ("与专业人员人工制图相比，本文系统的潜在优势主要体现在文本信息整理，重复性工具操作和质量反馈记录等环节。人工制图能够保持较高的专业判断和成果可控性，但需要手动阅读文本，整理节点，绘制路线，配置符号和组织版式，过程耗时且难以形成可复用的中间状态。本文系统将路线节点，事件说明，表达要求和评价结果记录为结构化状态，使同类路线图任务可以复用抽取模式，知识关系和工具链。不同方法在路线类专题图任务中的制图效果对比见表 5-1。", None),
    ("表 5-1 不同方法在路线类专题图任务中的制图效果对比", "table_performance"),
    ("与现有智能制图系统相比，本文系统更偏向特定专题图类型的知识约束制图。Smart-QGIS 等系统更关注自然语言到通用 GIS 工具的协议化调用。MapGPT 侧重于大语言模型与制图工具集成的地图生成。CartoAgent 强调多模态大模型和多智能体机制在地图风格迁移与评价中的应用。相比之下，本文系统将研究对象限定为路线类专题图，重点处理非结构化路线文本中的时间，地点，事件和路线关系，并将其转化为路线段，节点标注，事件说明框和版式要素等可执行制图内容。不同系统的能力侧重点见表 5-2。", None),
    ("表 5-2 AutoCarto 智能制图系统与代表性智能制图系统能力对比", "table_systems"),
    ("当前研究仍存在三方面限制。第一，实验对象主要集中在路线类专题图，尚未覆盖统计专题图，分区设色图，流向图和多变量综合图等其他地图类型，因此方法的通用性仍需通过更多任务验证。第二，路线文本解析和事件地点匹配仍依赖模型对自然语言的理解，当文本存在地点别名，历史地名，模糊方位或缺失坐标时，系统仍需要人工校正或外部地名知识库支持。第三，地图质量评价中的视觉判断仍具有一定主观性，尤其是标注遮挡，版式均衡和符号层次等问题，需要进一步结合人工评价，规则约束和多模态模型输出建立更稳定的评价标准。", None),
    ("因此，本文方法更适合作为路线类专题图智能制图的原型框架，而不是已经完全成熟的通用制图系统。后续工作应补充不同复杂度路线任务的定量实验，完善任务完成率，工具调用成功率，质量评分和修订效果等结果。同时，应引入更丰富的地名知识，历史事件知识和制图规则，提升系统对复杂文本和复杂路线关系的处理能力，并进一步完善地图质量评价标准。", None),
]


def set_cell_text(cell, text, red=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(9)
    if red:
        run.font.color.rgb = RED


def style_table(table: Table):
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)


def insert_paragraph_after(paragraph: Paragraph, text="", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style and style not in {"fig_placeholder", "red"}:
        new_para.style = style
    if text:
        run = new_para.add_run(text)
        if style == "red":
            run.font.color.rgb = RED
    return new_para


def insert_table_after(doc: Document, anchor: Paragraph, kind: str) -> Paragraph:
    if kind == "table_metrics":
        rows = [
            ["指标类别", "评价指标", "含义"],
            ["过程指标", "任务完成率", "成功生成目标地图成果的任务比例"],
            ["过程指标", "工具调用成功率", "工具调用返回成功状态的比例"],
            ["过程指标", "人工干预次数", "任务执行过程中需要人工修正的次数"],
            ["过程指标", "制图耗时", "从任务输入到成果导出的总时间"],
            ["成果指标", "节点完整性", "关键路线节点是否完整表达与标注"],
            ["成果指标", "路线连贯性", "路线是否按知识图谱顺序正确连接"],
            ["成果指标", "事件标注准确性", "事件说明是否与对应空间节点匹配"],
            ["成果指标", "版式完整性", "标题，图例，比例尺，指北针和出图范围是否完整"],
        ]
    elif kind == "table_performance":
        rows = [
            ["方法", "任务完成率/%", "平均人工干预次数/次", "平均工具调用错误次数/次", "平均耗时/min", "平均质量评分"],
            ["专业人员人工制图", "待补充", "待补充", "待补充", "待补充", "待补充"],
            ["通用大语言模型生成 PyQGIS 脚本", "待补充", "待补充", "待补充", "待补充", "待补充"],
            ["本文系统", "待补充", "待补充", "待补充", "待补充", "待补充"],
        ]
    else:
        rows = [
            ["系统", "核心定位", "非结构化文本处理", "路线语义约束", "质量反馈闭环", "本地部署支持"],
            ["Smart-QGIS[3]", "通用空间数据处理与制图智能体", "部分支持", "不支持", "不支持", "支持"],
            ["MapGPT[4]", "大语言模型与制图工具集成的地图生成框架", "部分支持", "不支持", "不支持", "部分支持"],
            ["CartoAgent[8]", "地图风格迁移与评价的多智能体制图框架", "部分支持", "不支持", "部分支持", "部分支持"],
            ["本文系统", "面向路线类专题图的闭环制图智能体", "支持", "支持", "支持", "支持"],
        ]

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    style_table(table)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            set_cell_text(table.cell(i, j), value, red=value == "待补充")
    anchor._p.addnext(table._tbl)
    return insert_paragraph_after(Paragraph(table._tbl, anchor._parent), "")


def remove_element(el):
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def paragraph_text(el):
    return "".join(t.text for t in el.iter() if t.tag.endswith("}t") and t.text).strip()


def replace_text_in_runs(doc: Document):
    replacements = {
        "AutoCarto AI": "AutoCarto 智能制图系统",
        "生成—评价—修订": "生成，评价与修订",
        "文本解析—知识组织—地图生成—质量反馈": "文本解析，知识组织，地图生成和质量反馈",
        "遵义—桐梓—娄山关—赤水": "遵义，桐梓，娄山关和赤水",
    }
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for old, new in replacements.items():
                if old in run.text:
                    run.text = run.text.replace(old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        for old, new in replacements.items():
                            if old in run.text:
                                run.text = run.text.replace(old, new)


def main():
    doc = Document(SRC)
    body = doc.element.body
    children = list(body)
    start = end = None
    for i, el in enumerate(children):
        if el.tag.endswith("}p"):
            text = paragraph_text(el)
            if text == "3 系统架构与方法":
                start = i
            elif text == "6 结论与展望":
                end = i
                break
    if start is None or end is None:
        raise RuntimeError("Cannot locate section range from chapter 3 to chapter 6")

    for el in children[end - 1 : start - 1 : -1]:
        remove_element(el)

    anchor_el = OxmlElement("w:p")
    children[end].addprevious(anchor_el)
    anchor = Paragraph(anchor_el, doc._body)

    current = anchor
    for text, style in CONTENT:
        if style and style.startswith("table_"):
            current = insert_paragraph_after(current, text)
            current.alignment = WD_ALIGN_PARAGRAPH.CENTER
            current = insert_table_after(doc, current, style)
        elif style == "fig_placeholder":
            current = insert_paragraph_after(current, text)
            current.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = current.add_run("（图件待补充）")
            run.font.color.rgb = RED
        else:
            current = insert_paragraph_after(current, text, style)

    remove_element(anchor._p)
    replace_text_in_runs(doc)

    if OUT.exists():
        OUT.unlink()
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
