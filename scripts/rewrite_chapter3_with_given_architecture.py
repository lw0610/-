from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zw/研究生/QGIS-agent/Eino")
SRC = Path("/Users/zw/Downloads/AutoCarto_专题图制图智能体论文-修订模式.docx")
OUT = ROOT / "AutoCarto_专题图制图智能体论文_第三章按总体架构图优化版.docx"


SECTION3 = [
    ("3 系统架构与方法", "Heading 1"),
    (
        "本文系统面向路线类专题图的自动制图任务，重点解决自然语言路线叙事难以直接转化为 GIS 制图操作的问题。与通用问答系统不同，路线类专题图制图同时涉及文本语义理解，空间数据组织，路线关系构建，专题符号表达，版式输出和成果校核等环节。若直接让大语言模型生成 PyQGIS 脚本，模型需要在一次输出中同时判断任务意图，构造数据路径，调用 QGIS 接口并完成版式设计，容易受到字段名称，图层状态，本地路径和执行环境的影响。参考 Smart-QGIS 等智能 GIS 系统关于自然语言驱动 GIS 工具调用的研究思路[3]，本文将大语言模型定位为任务理解和工具规划中心，将 QGIS 定位为确定性制图执行环境，并通过后端调度，工具契约和本地技能库连接二者。",
        None,
    ),
    (
        "基于这一思路，本文系统采用五层应用架构，如图 3-1 所示。自上而下分别为用户交互层，后端调度层，多智能体规划层，技能与工具层以及 GIS 制图执行层。该架构的核心不是简单堆叠软件模块，而是将用户需求，智能体规划，工具调用和 QGIS 执行组织为可传递的任务状态。用户输入经后端调度层封装为任务请求后，进入多智能体规划层进行意图解析，数据推理，专题表达，版式优化和质量评价。规划层生成的不是自然语言说明，而是受工具契约约束的调用序列。技能与工具层负责将调用序列转换为可执行操作，GIS 制图执行层则在 QGIS 本地实例中完成图层处理，地图渲染，布局生成和成果导出。执行状态和中间结果继续回传给规划层，使系统能够支持多轮任务推进和局部修订。",
        None,
    ),
    ("图 3-1 AutoCarto 智能制图系统应用总体架构", "figure"),
    ("3.1 应用总体架构", "Heading 2"),
    (
        "用户交互层为制图人员提供自然语言需求输入，文本，表格和业务说明接入，以及制图结果展示和成果预览。该层强调交互入口和成果回看，不直接承担制图推理。用户可以通过桌面端或 Web 界面提交路线类专题图制图需求，也可以查看地图，PDF 和布局工程等成果。交互层的设计保证了智能体任务能够以普通制图需求的形式被触发，同时使生成结果能够被用户及时检查。",
        None,
    ),
    (
        "后端调度层负责把用户请求转化为可执行任务。该层接收 WebSocket 或 HTTP 请求，完成请求接入与鉴权，并维护任务队列和状态机。制图任务在执行过程中往往需要多轮工具调用，单次模型回复并不能完成全部操作，因此后端调度层需要持续记录任务状态，执行日志和中间结果。执行状态回传机制使智能体能够根据 QGIS 返回结果调整后续规划，也使用户界面能够展示任务进度和异常信息。",
        None,
    ),
    (
        "多智能体规划层是自然语言制图需求转化为工具调用序列的关键环节。中央协调智能体负责识别用户请求中的工具调用意图，组织可用工具描述，并维护规划过程中的中间状态。意图解析智能体判断制图目标，路线范围和成果形式。数据推理智能体检查图层，字段，坐标和工程状态。专题表达智能体负责将路线节点，路线段和事件说明转化为专题符号表达。版式优化智能体处理标题，图例，比例尺，说明框和导出布局。质量评价智能体检查地图成果是否满足路线连贯性，节点完整性，图例一致性和版式可读性。上述智能体划分对应的是制图知识类型，而不是普通软件功能模块。",
        None,
    ),
    ("3.2 智能体规划与技能工具组织", "Heading 2"),
    (
        "本文系统采用模型规划和工具执行相分离的运行方式。大语言模型主要负责判断当前任务需要调用的工具及其参数，后端调度层负责解析工具调用，检查参数完整性，并将请求路由到对应工具。这样处理的原因在于，路线类专题图制图依赖大量外部状态，例如当前 QGIS 工程是否已经加载底图，路线节点图层是否存在，坐标字段是否完整，布局对象是否创建，导出路径是否有效。模型本身无法稳定维护这些外部状态，因此系统必须依靠工具执行结果和 QGIS 状态回传来校正规划过程。",
        None,
    ),
    (
        "技能与工具层为智能体规划提供约束。制图技能库保存与分类，符号，标注，图例和版式相关的领域知识。工具契约规定每个工具的输入，输出和约束，避免模型生成无法执行的任意命令。执行适配器负责参数检查和指令转换，将智能体产生的工具调用转化为 QGIS 侧能够执行的操作。项目中的 Skills 静态本地知识库在智能体创建时注入，使模型能够获得路线类专题图制图所需的规则，步骤和注意事项。工具注册机制则保证模型只能在已声明能力范围内进行规划。",
        None,
    ),
    (
        "在项目实现中，中央协调智能体并不直接访问 QGIS 内部对象，而是通过工具层获得有限且明确的操作能力。可调用工具包括数据加载，图层检查，路线段构建，节点标注，符号配置，版式创建和成果导出等。每个工具的返回结果都会被记录为任务状态的一部分，并反馈给智能体继续决策。例如，数据推理阶段若发现节点坐标字段缺失，系统不会继续执行路线生成，而是将字段缺失作为状态返回。质量评价阶段若发现事件说明框遮挡路线，系统会将该问题转化为版式修订请求，再调用相应工具调整说明框位置。由此，模型输出不再是一次性脚本，而是可验证，可回退和可修订的中间决策。",
        None,
    ),
    ("3.3 QGIS 本地执行与闭环反馈", "Heading 2"),
    (
        "GIS 制图执行层以 QGIS 本地实例作为真实地图生成环境。QGIS 本地实例负责工程，图层和布局管理，PyQt 信号用于将后端工具请求安全传递到 QGIS 主线程，QGIS 主线程再调用 PyQGIS 核心接口完成地图渲染和布局输出。这种执行方式符合桌面 GIS 插件的运行约束，也避免了模型直接操作 QGIS 对象带来的线程安全和状态不可控问题。最终地图成果通过 QGIS 地图画布和打印布局导出，支持地图画布预览，PDF 成果和布局工程保存。",
        None,
    ),
    (
        "该执行层的作用不仅是完成工具调用，还承担结果证据回传。每次工具执行后，系统都会获得图层状态，工具日志，中间结果或导出成果。后端调度层将这些状态返回给多智能体规划层，使后续规划能够基于真实执行结果继续推进。若数据加载失败，系统可以回到数据推理阶段重新检查路径或字段。若地图导出成功但质量评价发现标注遮挡，系统可以回到版式优化阶段进行局部修订。由此，系统形成会话隔离，QGIS 实例隔离和工程隔离基础上的持续交互机制。",
        None,
    ),
    (
        "从路线类专题图制图角度看，闭环反馈使系统从能否生成地图转向生成地图是否符合制图意图。文本需求提供任务目标，技能与工具层提供可执行能力，QGIS 执行层提供真实地图结果，质量评价智能体则将结果中的问题反馈给规划层。该机制使地图生成不再以一次导出为终点，而是能够围绕路线连贯性，节点完整性，事件注记准确性，图例一致性和版式清晰度进行迭代修订。这也是本文系统区别于普通自然语言 GIS 工具调用框架的关键所在。",
        None,
    ),
]


def insert_paragraph_after(paragraph: Paragraph, text="", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style and style != "figure":
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def delete_paragraph(paragraph: Paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def insert_existing_image_after(anchor: Paragraph, image_paragraph_xml, caption: str) -> Paragraph:
    image_xml = deepcopy(image_paragraph_xml)
    anchor._p.addnext(image_xml)
    img_para = Paragraph(image_xml, anchor._parent)
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = insert_paragraph_after(img_para, caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.font.size = Pt(9)
    return cap


def find_chapter3_range_and_image(doc: Document):
    start = end = None
    image_xml = None
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text == "3 系统架构与方法":
            start = i
        elif start is not None and end is None and text.startswith("4 "):
            end = i
            break
        if start is not None and image_xml is None:
            if any(r._r.xpath(".//w:drawing") for r in para.runs):
                image_xml = deepcopy(para._p)
    if start is None or end is None:
        raise RuntimeError("Cannot locate chapter 3 range")
    if image_xml is None:
        raise RuntimeError("Cannot locate architecture image in chapter 3")
    return start, end, image_xml


def main():
    doc = Document(SRC)
    start, end, image_xml = find_chapter3_range_and_image(doc)

    for idx in range(end - 1, start - 1, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start - 1]
    for text, style in SECTION3:
        if style == "figure":
            anchor = insert_existing_image_after(anchor, image_xml, text)
        else:
            anchor = insert_paragraph_after(anchor, text, style)

    if OUT.exists():
        OUT.unlink()
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
