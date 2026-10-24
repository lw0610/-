from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, RGBColor, Pt
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


ROOT = Path("/Users/zw/研究生/QGIS-agent/Eino")
SRC = Path("/Users/zw/Downloads/AutoCarto_专题图制图智能体论文-修订模式.docx")
OUT = ROOT / "AutoCarto_专题图制图智能体论文_第三章期刊风格优化版.docx"
FIG = ROOT / "docs" / "paper_figures" / "chapter3_journal_method_figure.png"


def font(size: int):
    for path in [
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/Supplemental/Songti.ttc",
    ]:
        if Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def text_size(draw, text, fnt):
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def center(draw, x1, y, x2, text, size=22, fill="#172033"):
    f = font(size)
    tw, _ = text_size(draw, text, f)
    draw.text((x1 + (x2 - x1 - tw) / 2, y), text, font=f, fill=fill)


def wrap(draw, text, size, width):
    f = font(size)
    lines, cur = [], ""
    for ch in text:
        trial = cur + ch
        if text_size(draw, trial, f)[0] <= width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = ch
    if cur:
        lines.append(cur)
    return lines


def box(draw, xy, title, body="", fill="#FFFFFF", outline="#CBD5E1", title_size=21, body_size=16):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=2)
    center(draw, x1 + 8, y1 + 16, x2 - 8, title, title_size)
    yy = y1 + 50
    for line in wrap(draw, body, body_size, x2 - x1 - 28):
        center(draw, x1 + 12, yy, x2 - 12, line, body_size, "#667085")
        yy += body_size + 8


def arrow(draw, p1, p2, color="#667085", width=3):
    import math
    draw.line([p1, p2], fill=color, width=width)
    a = math.atan2(p2[1] - p1[1], p2[0] - p1[0])
    s = 12
    p3 = (p2[0] - s * math.cos(a - math.pi / 6), p2[1] - s * math.sin(a - math.pi / 6))
    p4 = (p2[0] - s * math.cos(a + math.pi / 6), p2[1] - s * math.sin(a + math.pi / 6))
    draw.polygon([p2, p3, p4], fill=color)


def panel_label(draw, x, y, letter, title):
    draw.text((x, y), letter, font=font(30), fill="#172033")
    draw.text((x + 42, y + 3), title, font=font(24), fill="#172033")


def make_figure():
    FIG.parent.mkdir(parents=True, exist_ok=True)
    w, h = 2400, 1450
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((80, 45), "知识约束的路线类专题图智能制图方法", font=font(40), fill="#172033")
    d.text((80, 96), "从路线叙事到地图成果的语义表示、知识约束和反馈证据", font=font(22), fill="#667085")

    # a. Route narrative to semantic representation.
    panel_label(d, 85, 170, "a", "路线叙事的结构化表示")
    box(d, (95, 245, 780, 430), "路线叙事文本", "从 A 地出发，经 B 地转移，于 C 地发生战斗，随后抵达 D 地。", "#F8FAFC")
    arrow(d, (438, 430), (438, 515))
    box(d, (95, 515, 780, 760), "制图语义状态", "行动主体，时间线索，地点实体，路线节点，历史事件，节点顺序，表达要求", "#EAF6F0", "#A8CFC0")
    box(d, (95, 845, 780, 1190), "结构化输出示例", "RouteNode: A, B, C, D\nEvent: battle at C\nOrder: A, B, C, D\nMapNeed: main route, node label, event callout", "#FFFFFF")

    # b. Knowledge graph.
    panel_label(d, 910, 170, "b", "路线制图知识组织")
    nodes = {
        "制图任务": (1320, 650, 110, 48, "#E8F2F9"),
        "路线节点": (1075, 470, 96, 40, "#EAF6F0"),
        "路线段": (1565, 470, 85, 40, "#EAF6F0"),
        "历史事件": (1060, 830, 96, 40, "#FFF4CF"),
        "表达方法": (1590, 830, 96, 40, "#F0ECFA"),
        "地图要素": (1320, 330, 96, 40, "#F0ECFA"),
        "质量约束": (1830, 650, 96, 40, "#FCECEF"),
        "输出成果": (820, 650, 96, 40, "#E5E7EB"),
    }
    edges = [
        ("制图任务", "路线节点", "包含"),
        ("路线节点", "路线段", "连接"),
        ("路线节点", "历史事件", "承载"),
        ("路线段", "表达方法", "采用"),
        ("表达方法", "地图要素", "生成"),
        ("制图任务", "质量约束", "受控"),
        ("制图任务", "输出成果", "产生"),
        ("地图要素", "制图任务", "服务"),
    ]
    for a, b, lab in edges:
        x1, y1, *_ = nodes[a]
        x2, y2, *_ = nodes[b]
        d.line([(x1, y1), (x2, y2)], fill="#AEBBCC", width=3)
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        box(d, (mx - 42, my - 20, mx + 42, my + 20), lab, "", "#FFFFFF", "#DDE4EE", 15, 10)
    for name, (cx, cy, rx, ry, fill) in nodes.items():
        d.ellipse((cx - rx, cy - ry, cx + rx, cy + ry), fill=fill, outline="#7F8EA3", width=3)
        center(d, cx - rx, cy - 13, cx + rx, name, 19)
    box(d, (895, 1060, 1780, 1180), "知识约束作用", "限定抽取字段，校验路线顺序，约束工具输入，并为质量评价提供参照。", "#F8FAFC")

    # c. Execution and evidence.
    panel_label(d, 85, 1245, "c", "制图执行与反馈证据")
    x0, y0 = 390, 1275
    for i, (name, fill) in enumerate([
        ("工具计划", "#E8F2F9"),
        ("QGIS 执行状态", "#EAF6F0"),
        ("地图图像", "#F0ECFA"),
        ("评价结果", "#FCECEF"),
    ]):
        x = x0 + i * 400
        box(d, (x, y0, x + 285, y0 + 95), name, "", fill)
        if i < 3:
            arrow(d, (x + 285, y0 + 48), (x + 385, y0 + 48), "#667085", 3)
    d.text((1985, 1297), "反馈到路线表达和版式编排", font=font(20), fill="#667085")

    img.save(FIG, dpi=(300, 300))


SECTION3 = [
    ("3 系统架构与方法", "Heading 1"),
    ("本文提出一种面向路线类专题图的知识约束智能制图方法。该方法面向自然语言路线叙事难以直接转化为 GIS 制图操作，通用大模型缺少路线制图语义约束，自动生成地图缺少质量反馈等问题，将路线语义建模，知识图谱约束，多智能体协同，QGIS 工具执行和质量反馈机制组织为统一框架。与直接让大语言模型生成 PyQGIS 脚本不同，本文方法在工具执行前构造可检查的制图状态，使路线节点，历史事件，空间位置，表达方法，工具输入和评价结果能够在同一任务链中传递与修订。", None),
    ("图 3-1 知识约束的路线类专题图智能制图方法框架", "figure"),
    ("3.1 方法总体设计", "Heading 2"),
    ("路线类专题图智能制图任务可以定义为从自然语言路线叙事到可验证地图成果的转换过程。给定用户制图指令，历史文献或业务文本，系统需要识别其中的行动主体，时间线索，地点实体，路线节点，历史事件，节点顺序和成果表达要求，并结合已有空间数据与 QGIS 工程状态，生成包含路线线状符号，节点标注，事件说明，图例，比例尺，指北针和输出布局的专题地图。该任务的输入包括文本需求，空间数据，制图知识，模型配置和工具状态。输出包括结构化制图需求，路线制图知识图谱，工具调用序列，QGIS 图层与布局对象以及最终导出的地图成果。", None),
    ("AutoCarto 智能制图系统采用用户交互层，Go 后端调度层，多智能体层和 QGIS 执行层组成的四层架构。用户交互层负责接收自然语言制图需求，路线文本和制图参数，并展示任务状态，工具日志和地图成果。Go 后端调度层负责维护会话上下文，任务队列，模型配置，工具网关和运行日志，将前端输入转化为可调度的结构化请求。多智能体层负责完成路线意图解析，数据可用性推理，路线关系组织，专题表达设计，版式编排和质量评价。QGIS 执行层以 QGIS 和 PyQGIS 为本地制图环境，将工具调用转化为图层加载，几何处理，符号化，标注配置，布局生成和成果导出等操作。", None),
    ("该总体设计的关键在于区分语言推理，任务调度和 GIS 执行的边界。大语言模型负责理解制图意图，组织路线语义并选择合适工具。Go 后端负责运行时状态维护，工具调用转发和异常记录。QGIS 负责可验证的空间数据处理和地图渲染。通过这种分层设计，系统既避免了模型直接生成长脚本带来的不可控问题，也避免了 GIS 工具孤立执行而缺少语义约束的问题。", None),
    ("3.2 路线制图语义建模与知识组织", "Heading 2"),
    ("路线类专题图的核心信息通常以叙事方式出现，例如某一行动主体在特定时间从起点出发，经过若干地点，并在途中发生战斗，转移，驻留或会师等事件。此类文本不能直接作为 GIS 绘图参数使用，因为地点实体，路线顺序，事件归属和表达意图往往隐含在句子结构中。为此，本文将路线制图信息抽取定义为本体约束下的语义结构化过程，即将原始文本解析为可供知识图谱，智能体规划和工具执行共享的制图需求。", None),
    ("抽取模式包括行动主体，时间信息，地点实体，路线节点，历史事件，路线顺序，节点重要性，事件说明和成果表达要求等字段。地点实体记录名称，别名，地理线索和地理编码需求。路线节点记录节点序号，节点类型，关联地点和标注等级。历史事件记录事件名称，发生时间，关联节点和说明文本。表达要求记录路线线型，节点符号，说明框，图例和导出格式等需求。模型输出需要满足预定义字段集合和类型约束，并经过字段完整性，节点顺序一致性和事件地点匹配检查。", None),
    ("在语义抽取基础上，系统构建面向路线类专题图的制图本体与知识图谱。该知识图谱不是通用地理知识库，而是围绕路线图自动制图任务组织的领域知识结构，用于显式描述制图任务，路线节点，路线段，历史事件，空间位置，行动主体，表达方法，地图要素，质量约束和输出成果之间的关系。通过这些关系，系统能够在工具执行前明确哪些节点需要连接，哪些事件需要标注，哪些符号需要进入图例，以及哪些质量问题需要检查。", None),
    ("3.3 多智能体协同与制图反馈机制", "Heading 2"),
    ("路线类专题图制图链条较长，涉及文本理解，数据检查，路线组织，符号设计，版式编排和质量评价等多类知识。若由单一模型一次性完成，容易出现职责混淆，执行顺序错误和质量问题难以定位等情况。本文依据路线类专题图制图过程设置中央协调智能体，路线意图解析智能体，数据推理智能体，路线表达智能体，版式编排智能体和质量评价智能体，使不同知识类型由相应角色处理。", None),
    ("各智能体之间传递的是结构化制图需求，图谱关系，工具执行结果和评价状态，而不是仅依赖自由文本对话。中央协调智能体维护全局目标，任务状态和子任务顺序。路线意图解析智能体识别制图目标，行动主体，路线范围和成果形式。数据推理智能体检查路线节点，坐标字段，图层状态和 QGIS 工程环境。路线表达智能体依据图谱关系生成主路线，分支路线，多起点汇合路线和事件节点表达方案。版式编排智能体组织标题，图例，比例尺，指北针，节点标注和说明框。质量评价智能体根据知识图谱，QGIS 状态和地图图像输出问题定位与修订建议。", None),
    ("QGIS 制图执行采用工具契约而非自由脚本生成的方式组织。每个工具均限定输入参数，输出结果，适用条件和失败反馈，并由 Go 后端统一调度。工具链包括数据接入与检查，路线关系构建，路线表达，地图要素配置和成果导出等能力。工具执行过程中产生的图层状态，日志，布局对象和导出成果会被写入执行状态，供质量评价模块进一步检查。评价结果随后被转化为修订指令，反馈给路线表达智能体或版式编排智能体，使系统形成生成，评价与修订相结合的闭环制图过程。", None),
]


def insert_paragraph_after(paragraph: Paragraph, text="", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style and style != "figure":
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def delete_paragraph(paragraph: Paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def add_figure_after(paragraph: Paragraph, image_path: Path, caption: str) -> Paragraph:
    img_p = insert_paragraph_after(paragraph)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.paragraph_format.space_before = Pt(6)
    img_p.paragraph_format.space_after = Pt(2)
    img_p.add_run().add_picture(str(image_path), width=Cm(15.8))
    cap_p = insert_paragraph_after(img_p, caption)
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap_p.runs:
        run.font.size = Pt(9)
    return cap_p


def normalize_text(doc: Document):
    replacements = {
        "AutoCarto AI": "AutoCarto 智能制图系统",
        "生成—评价—修订": "生成，评价与修订",
        "文本信息抽取—图谱关系构建—智能体任务规划—QGIS 工具执行—地图质量评价—反馈修订": "文本信息抽取，图谱关系构建，智能体任务规划，QGIS 工具执行，地图质量评价和反馈修订",
        "遵义—桐梓—娄山关—赤水": "遵义，桐梓，娄山关和赤水",
    }
    for para in doc.paragraphs:
        for run in para.runs:
            for old, new in replacements.items():
                if old in run.text:
                    run.text = run.text.replace(old, new)


def main():
    make_figure()
    doc = Document(SRC)
    start = end = None
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text == "3 系统架构与方法":
            start = i
        elif start is not None and text.startswith("4 "):
            end = i
            break
    if start is None or end is None:
        raise RuntimeError("Cannot locate chapter 3")

    for idx in range(end - 1, start - 1, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start - 1]
    for text, kind in SECTION3:
        if kind == "figure":
            anchor = add_figure_after(anchor, FIG, text)
        else:
            anchor = insert_paragraph_after(anchor, text, kind)

    normalize_text(doc)
    if OUT.exists():
        OUT.unlink()
    doc.save(OUT)
    print(OUT)
    print(FIG)


if __name__ == "__main__":
    main()
