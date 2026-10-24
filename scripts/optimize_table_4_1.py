from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path("/Users/zw/研究生/QGIS-agent/Eino")
SRC = ROOT / "AutoCarto_专题图制图智能体论文_期刊风格结构版.docx"
OUT = ROOT / "AutoCarto_专题图制图智能体论文_表4-1优化版.docx"


ROWS = [
    ["评价维度", "指标", "计算或判定方式", "评价对象", "说明"],
    ["任务执行", "任务完成率", "成功完成任务数占任务总数的比例", "完整制图任务", "衡量系统从文本输入到成果导出的整体可用性"],
    ["任务执行", "工具调用成功率", "成功工具调用数占工具调用总数的比例", "工具调用日志", "反映工具契约和后端调度的稳定性"],
    ["任务执行", "人工干预次数", "人工修正或补充操作次数", "执行过程", "衡量系统对人工辅助的依赖程度"],
    ["任务执行", "制图耗时", "任务提交至成果导出的时间", "完整制图任务", "衡量自动制图过程的效率"],
    ["语义表达", "节点完整性", "正确表达节点数占应表达节点数的比例", "路线节点", "检查关键地点是否完整表达与标注"],
    ["语义表达", "路线连贯性", "正确连接路线段数占应连接路线段数的比例", "路线关系", "检查节点顺序和路线连接是否一致"],
    ["语义表达", "事件匹配准确性", "正确匹配事件数占事件总数的比例", "事件注记", "检查事件说明与空间节点是否对应"],
    ["地图表达", "图例一致性", "图例项与地图符号是否对应", "地图图例", "检查专题符号说明是否完整一致"],
    ["地图表达", "版式完整性", "必要版式要素是否齐全", "地图布局", "检查标题，图例，比例尺，指北针和出图范围"],
]


def set_cell(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(8)
    run.font.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def build_new_table(doc, old_table):
    anchor = old_table._tbl
    new_table = doc.add_table(rows=0, cols=len(ROWS[0]))
    anchor.addnext(new_table._tbl)
    anchor.getparent().remove(anchor)
    return new_table


def optimize_table(table):
    widths = [1300, 1700, 2700, 1700, 3600]
    for i, row_data in enumerate(ROWS):
        cells = table.add_row().cells
        for j, text in enumerate(row_data):
            set_cell(cells[j], text, bold=(i == 0))
            set_cell_width(cells[j], widths[j])
            if i == 0:
                set_cell_shading(cells[j], "EAF0F6")
            elif i % 2 == 0:
                set_cell_shading(cells[j], "F8FAFC")

    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "宋体"


def main():
    doc = Document(SRC)
    if not doc.tables:
        raise RuntimeError("No table found")
    table = build_new_table(doc, doc.tables[0])
    optimize_table(table)
    if OUT.exists():
        OUT.unlink()
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
