from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SRC = Path("/Users/zw/Downloads/AutoCarto_专题图制图智能体论文-修订模式.docx")
OUT = Path("/Users/zw/研究生/QGIS-agent/Eino/AutoCarto_专题图制图智能体论文_局部优化版.docx")


ABSTRACT = (
    "摘要：路线类专题图需要从业务文本和自然语言需求中识别行动主体、经过地点、事件过程、路线关系和表达要求，"
    "并进一步完成数据组织、专题表达、版式编排和成果校核，是专题图智能制图中语义关联较强、任务流程较长的一类问题。"
    "现有大语言模型辅助GIS研究多关注通用工具调用、脚本生成或局部地图设计，对路线类专题图中的文本语义解析、路线知识约束、"
    "多环节制图协同和成果质量反馈支持不足。针对上述问题，本文提出一种知识约束的路线类专题图智能制图方法。"
    "该方法首先从制图相关文本中抽取路线节点、事件信息、空间范围、表达约束和成果形式等要素，并构建由制图任务、路线节点、"
    "历史事件、路线关系、表达方法和地图要素组成的语义模型及知识图谱；随后依据路线类专题图制图流程，设计文本理解、数据处理、"
    "路线表达、版式编排和质量评价等任务型智能体，通过技能体系与QGIS工具链完成地图生成；最后引入多模态质量评价机制，"
    "从主题一致性、路线连贯性、节点完整性、图例一致性、标注可读性和版式规范性等方面反馈修订建议，形成需求解析、知识组织、"
    "智能体协同、地图生成和质量反馈相结合的闭环流程。以长征路线图为典型对象的实验表明，本文方法能够支撑从文本需求到路线类专题图成果的自动生成，"
    "并为面向业务文本的专题图智能制图提供了一种可解释、可扩展的实现路径。"
)

KEYWORDS = "关键词：路线类专题图；制图语义解析；知识图谱；多智能体协同；质量反馈；QGIS"

INTRO = [
    "专题图是表达地理现象空间分布、专题属性关系和时空过程的重要地图形式，在国土空间规划、资源环境管理、公共服务分析、交通组织、历史地理研究和应急管理等领域具有广泛应用价值。与普通地理底图主要提供空间参照不同，专题图制图需要围绕特定主题选择空间对象，组织专题属性，确定表达方法，并配置标题、图例、注记、比例尺和版式等地图要素。因此，专题图制图不仅是空间数据可视化过程，也是融合业务语义理解、制图知识组织和地图表达设计的综合性工作。",
    "在实际应用中，专题图制图需求往往并不以标准化参数形式给出，而是分散在规划报告、统计说明、历史文献、业务材料和用户自然语言描述之中。路线类专题图尤其具有代表性。以行军路线图、历史迁徙图、交通规划图和应急救援路线图为例，制图人员需要从文本中识别行动主体、时间线索、经过地点、路线顺序、关键事件和表达要求，再将这些语义信息转化为节点、线段、注记、说明框、图例和版式成果。该过程对制图人员的专业知识和GIS软件操作能力要求较高，容易出现需求理解不充分、路线关系组织不清、符号与图例不一致、成果反复修改等问题。",
    "近年来，大语言模型在自然语言理解、任务规划和工具调用方面表现出较强能力，为GIS操作和地图制图的智能化提供了新的技术路径。已有研究从通用GIS智能体、地图生成工具集成、多智能体地图设计和多模态地图评价等方面进行了探索，证明自然语言驱动空间数据处理和地图制作具有可行性[1-4,8]。这些工作降低了GIS操作门槛，也说明多智能体机制能够在一定程度上组织复杂制图流程。",
    "然而，面向路线类专题图制图任务，现有研究仍存在不足。首先，通用GIS智能体通常更关注自然语言到工具调用或脚本生成的映射，对业务文本中隐含的制图主题、空间对象、路线顺序、事件地点和表达约束缺少系统化解析。其次，路线类专题图不仅需要生成地图图层，还需要根据图种特点组织路线符号、节点标注、事件说明、图例和版式要素，现有方法对这类制图知识约束与跨环节状态传递支持不足。最后，已有研究多以地图生成为主要目标，对成果是否符合路线关系、表达规范和用户需求缺少持续的质量评价与反馈修订机制。",
    "针对上述问题，本文提出一种知识约束的路线类专题图智能制图方法。该方法面向制图相关文本进行语义解析，提取路线节点、事件信息、空间范围、表达约束和成果形式等要素，并将其组织为面向路线类专题图制图的结构化知识。在此基础上，依据路线类专题图制图流程构建任务型多智能体协同机制，使文本理解、数据处理、路线表达、版式编排和质量反馈等环节能够围绕同一制图任务协同执行。进一步地，本文将常见制图操作组织为可调用的技能链和工具链，通过知识约束降低工具调用和参数生成的不确定性，并结合地图质量评价结果对符号表达、图例配置、标注可读性和版式完整性进行反馈优化。",
    "本文的主要贡献包括以下三个方面。第一，提出面向路线类专题图制图的语义解析与知识组织方法，将文本材料和自然语言需求中的路线节点、事件信息、路线关系和表达约束转化为可用于任务规划和工具调用的结构化制图知识。第二，设计面向路线类专题图制图流程的任务型多智能体协同机制，围绕需求理解、数据检查、路线表达、版式编排和质量反馈等环节组织制图任务，提高复杂路线图任务的自动化执行能力。第三，构建知识约束下的路线类专题图技能链与质量反馈机制，将地图生成过程与成果评价修订相结合，提高自动制图过程的稳定性、可解释性和地图成果完整性。",
    "本文以长征路线图作为典型实验对象，对所提出的方法进行验证。长征路线图兼具空间对象组织、时序关系表达、事件信息标注、路线符号配置和版式要素编排等特点，能够较好反映路线类专题图智能制图中需求解析、知识组织、任务协同和质量反馈等关键问题。实验从文本解析准确性、制图任务完成率、人工干预次数、工具调用稳定性和地图质量评分等方面开展评价，并与直接生成PyQGIS脚本的方法进行对比分析。"
]

RELATED = [
    ("2.1 大语言模型驱动的GIS工具调用", "Heading 2"),
    ("大语言模型具备自然语言理解、任务规划和代码生成能力，为GIS操作方式从人工菜单交互向自然语言驱动转变提供了基础。Autonomous GIS探索以大语言模型为推理核心的自动化GIS框架，用于组织空间数据获取、空间分析和可视化任务[1]；ChatGeoAI将自然语言请求转化为可在PyQGIS环境中执行的代码，支持公众用户开展地理空间分析[2]；Smart-QGIS通过协议化通信机制连接大语言模型与QGIS原生工具，验证了自然语言驱动空间数据处理和地图制作的可行性[3]。MapGPT进一步从地图要素控制角度出发，将大语言模型与制图工具结合，用于自然语言驱动的地图生成[4]。", None),
    ("上述研究的共同贡献在于打通了大语言模型与GIS平台或制图工具之间的调用链路，使模型能够根据用户意图调用数据加载、空间处理、图层符号化和地图输出等功能。然而，路线类专题图制图不仅要求工具能够被调用，还要求系统理解文本中的路线节点、事件地点、时间顺序和表达要求。若缺少面向图种的语义约束，模型即使能够调用工具，也可能出现节点顺序错误、事件注记缺失或图例版式不完整等问题。", None),
    ("2.2 制图知识组织与专题图表达约束", "Heading 2"),
    ("专题图制图是一项高度依赖专业知识和表达规则的工作。传统自动制图和地图设计研究通常通过专家规则、模板约束或算法模型支持分类分级、色彩选择、符号配置、注记安排和版式布局。随着知识图谱和大语言模型的发展，将领域知识显式组织为可查询、可约束和可推理的结构化知识，逐渐成为智能制图研究的重要方向。知识图谱能够为大语言模型提供实体、关系和约束，有助于减少自由生成带来的事实漂移和逻辑不一致[5]。在地图制图场景中，制图知识不仅包括空间对象、属性字段和空间关系，也包括图种类型、专题属性、分类依据、视觉变量、图例规则、标注策略和成果形式等内容。", None),
    ("已有研究已开始探索将制图知识、配色规则或领域经验嵌入智能体流程。例如，MapColor-Agent面向行政区划图配色任务，将色彩知识、用户需求和多智能体流程结合，实现交互式地图配色方案生成[7]；相关地图生成研究也开始讨论领域知识对制图行为决策的约束作用[6]。但对于路线类专题图而言，制图知识还需要描述行动主体、路线节点、事件地点、先后顺序、分支汇入和事件说明等复合关系。现有通用制图知识模型对这类路线语义和制图要素之间的映射关系支持仍不充分。", None),
    ("2.3 多智能体制图与质量反馈", "Heading 2"),
    ("多智能体方法通过任务分解和角色协作，为复杂制图任务自动化提供了新的组织方式。CartoAgent将地图制作划分为准备、设计和评价等阶段，通过多模态大模型与多智能体协作完成地图风格迁移和评价[8]；在线地图地形可视化智能设计研究围绕地形表达过程设置多个任务智能体，支持地形样式生成与修订[9]。这些研究表明，多智能体机制能够将复杂地图设计任务拆解为职责明确的子任务，并通过协同执行提高地图设计过程的组织性和可控性。", None),
    ("与此同时，多模态大模型在视觉语言任务和地图阅读分析方面的能力，为地图成果自动评价提供了新的可能[10-11]。在路线类专题图中，质量评价不仅要判断视觉效果，还要检查路线关系是否连贯、关键节点是否完整、事件说明是否与地点匹配、图例是否覆盖主要符号以及版式是否便于阅读。现有多智能体制图研究多聚焦于风格迁移、配色或地形表达等局部设计任务，对路线类专题图从需求理解、知识组织、工具执行到质量反馈的完整流程支持仍不充分。", None),
    ("总体来看，已有研究已验证大语言模型与智能体技术在GIS工具调用、地图生成、制图知识组织和地图评价中的应用潜力，但面向路线类专题图智能制图仍存在进一步研究空间。一是业务文本和自然语言需求中的路线制图语义解析不足；二是制图知识与智能体流程结合不够紧密，路线表达规则、地图要素组织和工具调用过程缺少统一的结构化约束；三是地图生成与质量反馈之间尚未形成稳定衔接，评价结果难以直接转化为可执行的制图修订操作。基于此，本文研究知识约束下的路线类专题图语义解析、任务型多智能体协同和地图质量反馈机制，以支撑从文本需求到专题图成果生成的自动化制图流程。", None),
]

CONCLUSION = [
    "本文面向路线类专题图制图中自然语言需求难以直接转化为地图成果的问题，提出了一种知识约束的专题图智能制图方法。该方法从业务文本和自然语言描述中提取路线节点、事件信息、路线关系、表达约束和成果形式等信息，并将其组织为结构化制图知识；在此基础上，结合路线类专题图制图流程，构建由需求理解、数据处理、路线表达、版式编排和质量反馈等环节组成的多智能体协同制图过程。与直接生成脚本或调用工具的方式相比，该方法在工具执行前引入制图知识约束，使制图任务的理解、分解和执行具有更明确的依据。",
    "以长征路线图为典型任务的实验表明，本文方法能够将文本中的路线节点、事件信息、路线关系和表达要求转化为可执行的制图过程，并生成包含路线、节点、标注、图例和版式要素的专题图成果。与直接生成PyQGIS脚本的方法相比，知识约束和任务分解能够减少大语言模型自由生成带来的不确定性，在复杂制图信息组织、地图要素配置和成果输出稳定性方面表现出更好的适应性。该结果说明，面向路线类专题图构建结构化知识和任务型智能体协同机制，是提高智能制图可控性和成果完整性的一种可行路径。",
    "后续研究将从方法能力、应用场景和系统扩展三个方面继续完善。在方法能力方面，将进一步丰富路线类专题图知识表达，增强系统对复杂业务文本、隐含制图意图、多分支路线和多事件关系的理解能力；在应用场景方面，将围绕历史事件图、交通规划图、应急救援路线图和旅游导览图等任务扩展表达规则、制图技能和评价指标；在系统扩展方面，将加强与本地GIS工具链、空间数据服务和领域知识库的协同，并进一步提升质量评价结果向可执行修订指令转化的稳定性和可追踪性。"
]


def clear_paragraph(paragraph):
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_paragraph_text(paragraph, text, style=None):
    clear_paragraph(paragraph)
    if style:
        paragraph.style = style
    run = paragraph.add_run(text)
    return run


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def replace_range_after_heading(doc, heading_idx, delete_start, delete_end, items):
    # Delete bottom-up so indices remain valid.
    for idx in range(delete_end, delete_start - 1, -1):
        delete_paragraph(doc.paragraphs[idx])
    anchor = doc.paragraphs[heading_idx]
    for item in reversed(items):
        if isinstance(item, tuple):
            text, style = item
        else:
            text, style = item, None
        insert_paragraph_after(anchor, text, style=style)


def main():
    doc = Document(SRC)

    # Abstract paragraph currently contains both abstract and keywords.
    set_paragraph_text(doc.paragraphs[1], ABSTRACT)
    insert_paragraph_after(doc.paragraphs[1], KEYWORDS)

    # After inserting keywords, original indices shift by +1 after paragraph 1.
    # Intro: original 3-9 => current 4-10; heading original 2 => current 3.
    replace_range_after_heading(doc, heading_idx=3, delete_start=4, delete_end=10, items=INTRO)

    # Related work: find heading dynamically after intro replacement.
    heading_related = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "2 相关研究":
            heading_related = i
            break
    if heading_related is None:
        raise RuntimeError("Cannot locate 2 相关研究")

    # Delete until next top-level "3 " heading.
    next_h = None
    for i in range(heading_related + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].text.strip().startswith("3 "):
            next_h = i
            break
    if next_h is None:
        raise RuntimeError("Cannot locate section 3")
    replace_range_after_heading(doc, heading_related, heading_related + 1, next_h - 1, RELATED)

    # Conclusion: replace content after "6 结论与展望" until references.
    heading_conc = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "6 结论与展望":
            heading_conc = i
            break
    if heading_conc is None:
        raise RuntimeError("Cannot locate 6 结论与展望")
    ref_idx = None
    for i in range(heading_conc + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].text.strip() == "参考文献":
            ref_idx = i
            break
    if ref_idx is None:
        raise RuntimeError("Cannot locate 参考文献")
    replace_range_after_heading(doc, heading_conc, heading_conc + 1, ref_idx - 1, CONCLUSION)

    if OUT.exists():
        OUT.unlink()
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
