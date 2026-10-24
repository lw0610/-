from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


ROOT = Path("/Users/zw/研究生/QGIS-agent/Eino")
SRC = ROOT / "AutoCarto_专题图制图智能体论文_局部优化版_clean.docx"
OUT = ROOT / "AutoCarto_专题图制图智能体论文_方法章节优化版.docx"
FIG = ROOT / "docs" / "paper_figures" / "method_framework_route_map.png"


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/Supplemental/Songti.ttc",
    ]
    for p in candidates:
        if Path(p).exists():
            return ImageFont.truetype(p, size=size)
    return ImageFont.load_default()


def text_size(draw, text, f):
    b = draw.textbbox((0, 0), text, font=f)
    return b[2] - b[0], b[3] - b[1]


def wrap(draw, text, f, width):
    lines, cur = [], ""
    for ch in text:
        trial = cur + ch
        if text_size(draw, trial, f)[0] <= width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = ch
    if cur:
        lines.append(cur)
    return lines


def box(draw, xy, title, body, fill, outline="#64748B"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=2)
    ft = font(25)
    fb = font(19)
    tw, th = text_size(draw, title, ft)
    draw.text((x1 + (x2 - x1 - tw) / 2, y1 + 18), title, fill="#0F172A", font=ft)
    y = y1 + 62
    for line in wrap(draw, body, fb, x2 - x1 - 36):
        lw, _ = text_size(draw, line, fb)
        draw.text((x1 + (x2 - x1 - lw) / 2, y), line, fill="#334155", font=fb)
        y += 28


def arrow(draw, start, end, color="#334155", width=4):
    import math
    draw.line([start, end], fill=color, width=width)
    ang = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 16
    p1 = (end[0] - size * math.cos(ang - math.pi / 6), end[1] - size * math.sin(ang - math.pi / 6))
    p2 = (end[0] - size * math.cos(ang + math.pi / 6), end[1] - size * math.sin(ang + math.pi / 6))
    draw.polygon([end, p1, p2], fill=color)


def polyline_arrow(draw, points, color="#334155", width=4):
    import math
    draw.line(points, fill=color, width=width, joint="curve")
    start, end = points[-2], points[-1]
    ang = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 16
    p1 = (end[0] - size * math.cos(ang - math.pi / 6), end[1] - size * math.sin(ang - math.pi / 6))
    p2 = (end[0] - size * math.cos(ang + math.pi / 6), end[1] - size * math.sin(ang + math.pi / 6))
    draw.polygon([end, p1, p2], fill=color)


def make_figure():
    FIG.parent.mkdir(parents=True, exist_ok=True)
    W, H = 2200, 1320
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)

    def panel(x1, y1, x2, y2, tag, title):
        d.rounded_rectangle((x1, y1, x2, y2), radius=18, fill="#FFFFFF", outline="#CBD5E1", width=2)
        d.text((x1 + 24, y1 + 18), tag, fill="#0F172A", font=font(30))
        d.text((x1 + 78, y1 + 20), title, fill="#1E293B", font=font(26))

    def centered_text(x1, y1, x2, text, f, color="#0F172A"):
        tw, _ = text_size(d, text, f)
        d.text((x1 + (x2 - x1 - tw) / 2, y1), text, fill=color, font=f)

    def layer(x1, y1, x2, y2, title, body, fill, left_label):
        d.rounded_rectangle((x1, y1, x2, y2), radius=12, fill=fill, outline="#94A3B8", width=1)
        d.text((x1 + 18, y1 + 20), title, fill="#0F172A", font=font(23))
        body_f = font(18)
        yy = y1 + 57
        for line in wrap(d, body, body_f, x2 - x1 - 40):
            centered_text(x1 + 12, yy, x2 - 12, line, body_f, "#475569")
            yy += 25
        d.rounded_rectangle((x1 - 145, y1 + 10, x1 - 20, y2 - 10), radius=18, fill="#F8FAFC", outline="#CBD5E1", width=1)
        centered_text(x1 - 145, y1 + 34, x1 - 20, left_label, font(18), "#334155")

    def ellipse_node(cx, cy, rx, ry, title, fill="#F8FAFC", outline="#64748B"):
        d.ellipse((cx - rx, cy - ry, cx + rx, cy + ry), fill=fill, outline=outline, width=2)
        centered_text(cx - rx, cy - 14, cx + rx, title, font(21), "#0F172A")

    def relation(start, end, label=None):
        d.line([start, end], fill="#64748B", width=2)
        if label:
            mx, my = (start[0] + end[0]) / 2, (start[1] + end[1]) / 2
            d.rounded_rectangle((mx - 54, my - 17, mx + 54, my + 17), radius=8, fill="white", outline="#E2E8F0", width=1)
            centered_text(mx - 54, my - 12, mx + 54, label, font(15), "#475569")

    def small_box(x1, y1, x2, y2, title, fill="#F8FAFC"):
        d.rounded_rectangle((x1, y1, x2, y2), radius=12, fill=fill, outline="#CBD5E1", width=1)
        centered_text(x1, y1 + 22, x2, title, font(21), "#0F172A")

    # Panel a: layered system architecture
    panel(60, 55, 1025, 745, "a", "知识约束的系统层级架构")
    layers = [
        ("制图需求层", "路线文本、历史材料、自然语言任务", "#E0F2FE", "输入"),
        ("语义建模层", "节点、事件、时间、地点与表达约束抽取", "#ECFDF5", "抽取"),
        ("知识约束层", "本体概念、图谱关系、字段与工具契约", "#FEF3C7", "约束"),
        ("智能体协同层", "意图解析、数据推理、路线表达、版式编排、质量评价", "#FCE7F3", "规划"),
        ("GIS执行层", "QGIS图层、符号、标注、布局与导出", "#EDE9FE", "执行"),
    ]
    y0 = 135
    for i, item in enumerate(layers):
        layer(245, y0 + i * 112, 935, y0 + i * 112 + 78, *item)
        if i < len(layers) - 1:
            arrow(d, (590, y0 + i * 112 + 78), (590, y0 + (i + 1) * 112 - 8), "#94A3B8", 2)
    d.rounded_rectangle((250, 687, 930, 725), radius=16, fill="#F8FAFC", outline="#CBD5E1", width=1)
    centered_text(250, 696, 930, "输出：路线类专题图、布局文件、导出成果与运行日志", font(17), "#475569")

    # Panel b: ontology and knowledge graph schema
    panel(1080, 55, 2140, 745, "b", "路线制图本体与知识图谱关系")
    nodes = {
        "制图任务": (1605, 355, 105, 45, "#E0F2FE"),
        "路线节点": (1360, 220, 92, 38, "#ECFDF5"),
        "路线段": (1848, 220, 82, 38, "#ECFDF5"),
        "历史事件": (1328, 485, 92, 38, "#FEF3C7"),
        "空间位置": (1605, 585, 92, 38, "#F8FAFC"),
        "表达方法": (1880, 485, 92, 38, "#FCE7F3"),
        "地图要素": (1605, 140, 92, 38, "#EDE9FE"),
        "质量约束": (1995, 355, 92, 38, "#FFE4E6"),
        "输出成果": (1215, 355, 92, 38, "#E5E7EB"),
    }
    edges = [
        ("制图任务", "路线节点", "包含"),
        ("路线节点", "路线段", "连接"),
        ("路线节点", "历史事件", "承载"),
        ("历史事件", "空间位置", "发生于"),
        ("路线段", "表达方法", "采用"),
        ("表达方法", "地图要素", "生成"),
        ("制图任务", "质量约束", "受控于"),
        ("制图任务", "输出成果", "产生"),
        ("地图要素", "制图任务", "服务于"),
    ]
    for a, b, label in edges:
        ax, ay, _, _, _ = nodes[a]
        bx, by, _, _, _ = nodes[b]
        relation((ax, ay), (bx, by), label)
    for title, (cx, cy, rx, ry, fill) in nodes.items():
        ellipse_node(cx, cy, rx, ry, title, fill)
    d.rounded_rectangle((1130, 650, 2095, 707), radius=12, fill="#F8FAFC", outline="#E2E8F0", width=1)
    d.text((1160, 666), "图谱作用：将文本叙事转化为可检查的节点关系、表达约束和工具输入。", fill="#475569", font=font(20))

    # Panel c: evaluation-driven revision
    panel(60, 790, 2140, 1260, "c", "生成—评价—修订闭环与质量证据")
    cx, cy = 610, 1035
    loop = [
        (360, 920, 570, 990, "候选地图生成", "#E0F2FE"),
        (680, 920, 890, 990, "结构化检查", "#FEF3C7"),
        (680, 1080, 890, 1150, "视觉质量评价", "#FCE7F3"),
        (360, 1080, 570, 1150, "局部修订", "#ECFDF5"),
    ]
    for x1, y1, x2, y2, title, fill in loop:
        small_box(x1, y1, x2, y2, title, fill)
    arrow(d, (570, 955), (680, 955), "#64748B", 3)
    arrow(d, (785, 990), (785, 1080), "#64748B", 3)
    arrow(d, (680, 1115), (570, 1115), "#64748B", 3)
    arrow(d, (465, 1080), (465, 990), "#64748B", 3)
    d.ellipse((535, 995, 715, 1075), fill="#FFFFFF", outline="#94A3B8", width=2)
    centered_text(535, 1022, 715, "知识约束", font(21), "#0F172A")

    # Evidence matrix
    x0, y0 = 1085, 890
    d.text((x0, y0 - 45), "质量评价证据矩阵", fill="#1E293B", font=font(25))
    cols = ["知识图谱", "QGIS状态", "地图图像"]
    rows = ["路线连贯性", "节点完整性", "符号层次", "版式可读性"]
    cell_w, cell_h = 235, 58
    d.rounded_rectangle((x0, y0, x0 + 4 * cell_w, y0 + 5 * cell_h), radius=12, fill="#FFFFFF", outline="#CBD5E1", width=1)
    for j, col in enumerate(["评价维度"] + cols):
        d.rectangle((x0 + j * cell_w, y0, x0 + (j + 1) * cell_w, y0 + cell_h), fill="#F1F5F9", outline="#CBD5E1")
        centered_text(x0 + j * cell_w, y0 + 17, x0 + (j + 1) * cell_w, col, font(19), "#0F172A")
    marks = [
        ["●", "●", "○"],
        ["●", "●", "○"],
        ["○", "●", "●"],
        ["○", "●", "●"],
    ]
    for i, row in enumerate(rows):
        yy = y0 + (i + 1) * cell_h
        d.rectangle((x0, yy, x0 + cell_w, yy + cell_h), fill="#FFFFFF", outline="#CBD5E1")
        centered_text(x0, yy + 17, x0 + cell_w, row, font(18), "#334155")
        for j, mark in enumerate(marks[i]):
            xx = x0 + (j + 1) * cell_w
            d.rectangle((xx, yy, xx + cell_w, yy + cell_h), fill="#FFFFFF", outline="#CBD5E1")
            centered_text(xx, yy + 14, xx + cell_w, mark, font(24), "#0F766E" if mark == "●" else "#94A3B8")
    d.text((x0, y0 + 5 * cell_h + 24), "● 主要证据源    ○ 辅助证据源", fill="#64748B", font=font(18))

    img.save(FIG, dpi=(300, 300))


SECTION3 = [
    ("为解决路线类专题图中非结构化文本难以直接转化为可制图知识、通用GIS智能体缺少路线语义约束、地图生成结果缺少质量反馈等问题，本文构建了知识约束的路线类专题图智能制图方法。该方法以QGIS本地环境作为制图执行基础，以路线类专题图本体与知识图谱作为语义约束，以任务型多智能体机制组织制图流程，并通过质量评价结果驱动局部修订。整体流程不是将自然语言请求直接映射为脚本，而是在工具执行前形成可检查、可传递和可复用的路线制图知识。", None),
    ("图 3-1 知识约束的路线类专题图智能制图方法框架", "figure"),
    ("3.1 任务定义与总体框架", "Heading 2"),
    ("路线类专题图智能制图任务可形式化为从文本需求到地图成果的转换过程。给定用户自然语言指令、历史材料或业务文本，系统需要识别其中的行动主体、时间线索、地点实体、路线节点、事件信息、路线关系和表达要求，并结合已有空间数据和QGIS工程状态，生成包含路线、节点、注记、图例和版式要素的地图成果。该任务的输入包括文本需求、空间数据、制图知识和工具状态，输出包括结构化制图知识、工具调用序列、QGIS图层与布局对象以及最终导出的地图文件。", None),
    ("本文方法由五个环节组成。第一，制图语义解析环节从文本中抽取路线节点、历史事件、路线顺序和表达约束。第二，知识组织环节将抽取结果组织为路线制图知识图谱，显式描述制图任务、路线节点、历史事件、路线关系、表达方法和地图要素之间的关系。第三，任务规划环节由多智能体根据知识图谱和工程状态生成工具调用计划。第四，制图执行环节通过QGIS/PyQGIS工具链完成数据加载、路线生成、符号化、标注配置、版式生成和成果导出。第五，质量反馈环节结合图谱关系、图层状态和导出地图，对路线连贯性、节点完整性、事件标注准确性、图例一致性和版式可读性进行检查，并将问题转化为修订指令。", None),
    ("3.2 本体约束的路线制图信息抽取", "Heading 2"),
    ("路线类专题图的核心信息通常隐藏在自然语言叙事中，不能直接作为GIS绘图参数使用。为此，本文采用本体约束的信息抽取方式，将自由文本解析为结构化制图要素。抽取字段包括行动主体、时间信息、地点实体、路线节点、历史事件、路线顺序、事件说明、节点重要性和成果表达要求等。其中，地点实体记录名称、别名、地理线索和地理编码需求；路线节点记录节点序号、节点类型和标注等级；历史事件记录事件名称、发生时间、关联地点和说明文本。", None),
    ("与通用文本抽取不同，该环节的重点不是生成摘要，而是为后续制图流程提供稳定输入。模型输出需要满足预定义字段集合和类型约束，并经过字段完整性、节点顺序和事件地点匹配检查。这样可以降低模型自由生成造成的字段漂移和关系遗漏，使后续知识图谱构建、工具调用和质量评价具有统一的数据基础。", None),
    ("3.3 路线制图知识图谱构建", "Heading 2"),
    ("在完成信息抽取后，系统将路线节点、历史事件、路线关系、表达方法和地图要素组织为路线制图知识图谱。该知识图谱不是通用地理知识库，而是服务于路线类专题图自动制图的领域知识结构，用于描述文本叙事中的路线过程如何转化为可执行的地图表达。核心概念包括制图任务、路线节点、路线段、历史事件、空间位置、行动主体、表达方法、地图要素、质量约束和输出成果。", None),
    ("知识图谱中的典型关系包括：制图任务包含路线节点，路线节点连接相邻路线节点，路线段连接起点与终点，历史事件发生于路线节点，行动主体参与历史事件，路线类型对应线状符号，重要事件需要说明框，地图成果包含标题、图例、比例尺和指北针等地图要素。通过这些关系，系统能够将文本中的地点顺序、事件归属和表达要求显式化，使路线组织、节点标注和版式配置不再依赖模型对原始文本的临时理解。", None),
    ("3.4 任务型多智能体协同机制", "Heading 2"),
    ("路线类专题图制图链条较长，单一模型容易在需求理解、数据检查、路线组织、表达设计和质量评价之间发生职责混淆。本文依据路线图制图流程设置中央协调智能体、路线意图解析智能体、数据推理智能体、路线表达智能体、版式编排智能体和质量评价智能体。中央协调智能体维护任务状态和子任务顺序；路线意图解析智能体识别制图目标、行动主体、路线范围和输出形式；数据推理智能体检查节点坐标、图层状态、字段结构和QGIS工程环境；路线表达智能体依据图谱关系生成主路线、分支路线和事件节点表达方案；版式编排智能体组织标题、图例、比例尺、指北针、节点标注和说明框；质量评价智能体检查地图成果并生成修订建议。", None),
    ("各智能体之间传递的是结构化制图需求、知识图谱关系、工具执行结果和质量评价结果，而不是仅依赖自由文本对话。该设计使不同知识类型能够分离处理：文本语义由意图解析环节处理，数据状态由数据推理环节处理，路线关系由路线表达环节处理，地图成图由版式编排环节处理，质量问题由评价环节反馈。通过这种职责划分，系统能够在实验中追踪错误来源，并以工具调用成功率、人工干预次数和成果完整性等指标评价协同效果。", None),
    ("3.5 知识驱动的QGIS制图执行", "Heading 2"),
    ("路线图制图执行的关键是将知识图谱中的节点、事件和表达约束转化为QGIS中可执行的图层与布局操作。系统没有直接暴露全部QGIS工具，而是围绕路线类专题图制作过程组织制图技能和工具契约。每个工具均包含输入参数、输出结果、适用条件和失败反馈，用于约束模型生成可执行、可检查的工具调用。", None),
    ("工具链主要包括五类能力。第一类是数据接入与检查工具，用于加载路线节点、路线段、底图和行政边界数据，并检查坐标系统、字段名称和几何类型。第二类是路线关系构建工具，用于根据节点顺序、分支关系和多起点关系生成路线段。第三类是路线表达工具，用于生成平滑路线、箭头、主路线和分支路线符号，并处理路线叠压和层级关系。第四类是地图要素配置工具，用于生成节点标注、事件说明框、图例、标题、比例尺和指北针。第五类是成果导出工具，用于输出PDF、PNG或布局工程文件。", None),
    ("3.6 多模态质量评价与反馈优化", "Heading 2"),
    ("路线类专题图的质量不仅取决于工具是否成功执行，还取决于地图是否准确表达路线过程、节点是否完整、事件注记是否准确、路线层次是否清晰以及版式是否便于阅读。为此，系统将用户原始需求、路线知识图谱、QGIS图层与布局状态以及导出的地图图像共同作为质量评价输入，围绕路线连贯性、节点完整性、事件标注准确性、符号层次清晰度、图例一致性、标注可读性和版式完整性等维度开展检查。", None),
    ("评价过程采用规则检查与多模态模型判断相结合的方式。对于节点数量、图例项和地图要素完整性等可结构化检查内容，系统优先依据知识图谱和QGIS状态进行规则判断；对于视觉层次、标注遮挡和整体可读性等难以完全规则化的问题，系统将地图图像和评价提示输入多模态模型，生成问题描述和修改建议。评价结果进一步转化为修订指令，反馈给路线表达智能体或版式编排智能体，使系统从一次性地图生成扩展为“生成—评价—修订”的闭环制图流程。", None),
]


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def add_figure_after(paragraph, image_path, caption):
    img_p = insert_paragraph_after(paragraph)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.paragraph_format.space_before = Pt(6)
    img_p.paragraph_format.space_after = Pt(2)
    img_p.add_run().add_picture(str(image_path), width=Cm(15.6))
    cap_p = insert_paragraph_after(img_p, caption)
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.paragraph_format.space_after = Pt(8)
    for run in cap_p.runs:
        run.font.size = Pt(9)
    return cap_p


def main():
    make_figure()
    doc = Document(SRC)

    start = None
    end = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t == "3 系统架构与方法":
            start = i
        elif start is not None and t.startswith("4 "):
            end = i
            break
    if start is None or end is None:
        raise RuntimeError("Cannot locate section 3 range")

    # Delete all old section-3 body paragraphs and old figures, keep heading.
    for idx in range(end - 1, start, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start]
    for text, kind in SECTION3:
        if kind == "figure":
            anchor = add_figure_after(anchor, FIG, text)
        else:
            style = kind if kind else None
            anchor = insert_paragraph_after(anchor, text, style=style)

    if OUT.exists():
        OUT.unlink()
    doc.save(OUT)
    print(OUT)
    print(FIG)


if __name__ == "__main__":
    main()
