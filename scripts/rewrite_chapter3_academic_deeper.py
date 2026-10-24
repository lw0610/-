from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from rewrite_chapter3_journal_style import FIG, make_figure


ROOT = Path("/Users/zw/研究生/QGIS-agent/Eino")
SRC = ROOT / "AutoCarto_专题图制图智能体论文_第三章期刊风格优化版.docx"
OUT = ROOT / "AutoCarto_专题图制图智能体论文_第三章深化优化版.docx"


SECTION3 = [
    ("3 系统架构与方法", "Heading 1"),
    (
        "路线类专题图智能制图的关键并不只是让大语言模型调用 GIS 工具，而是将自然语言中的路线叙事转化为能够被 GIS 环境执行和被制图规则检验的中间表示。路线叙事通常同时包含行动主体，时间线索，地点实体，事件过程，路线顺序和表达意图。这些信息在文本中往往以叙述方式交织出现，既不是标准化空间数据，也不是可直接传入 PyQGIS 工具的参数。若直接由模型生成脚本，模型需要同时完成语义理解，空间关系组织，工具参数构造和版式设计，任何一个环节的字段缺失或关系误判都会传递到最终地图中。基于这一问题，本文将路线类专题图制图过程理解为一组可追踪状态之间的转换过程，并通过路线制图知识对模型规划和工具执行进行约束。",
        None,
    ),
    (
        "本文方法的核心思想是，在自然语言需求和 QGIS 制图操作之间引入路线制图语义状态。该状态既不同于原始文本，也不同于最终地图，而是由路线节点，历史事件，空间位置，表达方法，地图要素和质量约束共同组成的结构化对象。模型的作用不再是一次性生成完整地图脚本，而是围绕该结构化对象逐步完成信息抽取，关系组织，工具选择，地图表达和质量修订。这样，路线图制图过程可以从不可解释的文本到代码转换，转变为具有明确中间状态和检查依据的知识约束过程。图 3-1 给出了本文方法的总体框架。",
        None,
    ),
    ("图 3-1 知识约束的路线类专题图智能制图方法框架", "figure"),
    ("3.1 路线类专题图制图任务形式化", "Heading 2"),
    (
        "给定一段路线叙事文本和一组空间数据，路线类专题图制图任务可以表示为从文本语义到地图表达的映射。设输入文本为 D，空间数据集合为 S，QGIS 工程状态为 E，系统需要生成地图成果 M 以及对应的布局对象，图层状态和导出文件。与一般文本生成任务不同，该映射必须同时满足语义一致性和制图一致性。语义一致性要求地图中的路线节点，事件说明和路线顺序与原始文本保持对应。制图一致性要求地图符号，注记，图例和版式符合路线类专题图的表达习惯，并且能够在 QGIS 中稳定执行。",
        None,
    ),
    (
        "在这一任务定义下，本文将制图过程拆解为四类状态。第一类是文本语义状态，用于记录从文本中抽取的行动主体，时间线索，地点实体，路线节点，历史事件和成果表达要求。第二类是路线知识状态，用于记录节点之间的顺序关系，事件与地点的对应关系，路线段与表达方法之间的映射关系，以及地图要素与质量约束之间的关联。第三类是工具执行状态，用于记录 QGIS 工具调用，图层变化，样式配置，标注设置和布局导出结果。第四类是质量评价状态，用于记录节点缺失，路线断裂，事件错配，图例不一致和版式遮挡等问题。四类状态之间不是简单顺序关系，而是在制图过程中不断传递和更新。",
        None,
    ),
    (
        "这种形式化处理的意义在于，将路线图制图中的不确定性限制在可检查的状态转换中。文本解析阶段的不确定性体现在地点别名，路线顺序和事件归属等问题上。工具执行阶段的不确定性体现在图层路径，字段名称，坐标系统和布局参数等问题上。地图评价阶段的不确定性体现在视觉层次，标注遮挡和版式可读性等问题上。通过将这些问题分别映射到语义状态，知识状态，执行状态和评价状态，系统能够在不同阶段使用不同证据进行约束，而不是依赖模型对完整制图流程的一次性自由生成。",
        None,
    ),
    ("3.2 路线制图知识表示与约束传播", "Heading 2"),
    (
        "路线类专题图的知识表示需要同时覆盖文本语义和地图表达。本文构建的路线制图知识并非通用地理知识库，而是服务于路线图自动制图的任务型知识结构。其基本对象包括制图任务，行动主体，路线节点，路线段，历史事件，空间位置，表达方法，地图要素，质量约束和输出成果。制图任务用于限定地图类型和成果形式。路线节点用于表示文本中具有空间意义的关键地点。路线段用于表示相邻节点之间的连接关系。历史事件用于表示与路线节点相关的战斗，转移，驻留或会师等过程。表达方法用于描述线型，箭头，节点符号，标注和说明框等地图表达方式。质量约束用于描述路线连贯性，节点完整性，事件匹配和版式可读性等评价要求。",
        None,
    ),
    (
        "知识表示的关键不是罗列概念，而是建立概念之间的约束关系。路线节点之间的先后关系约束路线段生成，历史事件与空间位置之间的发生关系约束事件注记位置，路线类型与表达方法之间的映射关系约束线状符号配置，地图要素与输出成果之间的包含关系约束布局完整性，质量约束与图谱对象之间的对应关系约束评价依据。通过这些关系，系统能够将文本中的隐含路线过程转化为可查询的制图知识。后续工具调用并不直接读取原始文本，而是读取已经被知识结构组织过的节点，事件，关系和表达要求。",
        None,
    ),
    (
        "约束传播体现在三个层面。首先，在信息抽取层面，预定义的本体字段限制模型输出的结构，使抽取结果必须包含节点顺序，事件地点和表达要求等制图必要信息。其次，在任务规划层面，知识图谱中的节点连接关系和事件归属关系限制智能体的工具选择，使路线生成，节点标注和说明框配置具有明确依据。最后，在质量评价层面，图谱对象与地图对象之间的对应关系为检查提供参照，使系统能够判断地图中是否遗漏关键节点，路线是否按照预期顺序连接，事件注记是否与对应地点匹配，图例是否覆盖主要专题符号。由此，知识图谱不仅是信息存储结构，也是贯穿制图流程的约束机制。",
        None,
    ),
    ("3.3 知识约束下的智能体制图与反馈修订", "Heading 2"),
    (
        "在完成路线制图知识组织后，系统需要将知识状态转化为 QGIS 中可执行的制图操作。本文采用任务型多智能体协同机制来完成这一转换。与按照软件层级划分功能不同，智能体划分依据的是路线图制图中的知识类型。路线意图解析智能体处理文本需求与制图目标，数据推理智能体处理图层，字段，坐标和工程状态，路线表达智能体处理节点连接和路线符号，版式编排智能体处理标题，图例，比例尺，指北针，标注和说明框，质量评价智能体处理地图成果与知识状态之间的一致性。中央协调智能体维护任务状态，并根据中间结果决定下一步工具调用。",
        None,
    ),
    (
        "智能体之间传递的不是自由对话文本，而是结构化状态。路线表达智能体接收的是路线节点顺序，路线段关系和表达约束，而不是完整原始文本。版式编排智能体接收的是地图要素清单，图层状态和输出要求，而不是笼统的排版指令。质量评价智能体接收的是知识图谱，QGIS 状态和导出地图图像，而不是仅凭主观描述判断地图好坏。这样的状态传递方式使系统能够在工具执行失败或成果质量不足时定位问题来源，并将修订指令反馈到具体环节。",
        None,
    ),
    (
        "QGIS 执行环节采用工具契约控制模型输出。每个工具均限定输入参数，输出结果，适用条件和失败反馈。数据接入工具检查路径，坐标系统和字段结构。路线构建工具依据节点顺序生成路线段。表达配置工具设置线型，箭头，颜色，层级和节点符号。布局工具配置标题，图例，比例尺，指北针，标注和说明框。导出工具生成地图成果和布局文件。工具契约将复杂制图流程拆分为可检查的局部操作，降低了模型一次性生成长脚本时常见的语法错误，字段错配和环境假设问题。",
        None,
    ),
    (
        "反馈修订是本文方法区别于一次性自动制图流程的重要环节。地图导出后，系统并不立即结束任务，而是将导出的地图图像，QGIS 图层与布局状态，路线知识图谱和原始制图需求共同作为评价证据。对于节点数量，路线顺序和地图要素完整性等可结构化检查内容，系统依据知识图谱和工程状态进行规则判断。对于标注遮挡，符号层次和版式可读性等难以完全规则化的问题，系统结合地图图像进行多模态评价。评价结果进一步转化为局部修订动作，例如补充关键节点标注，调整主路线与分支路线的层级关系，重新计算说明框位置，或补充图例项。通过这一机制，路线类专题图制图由单次生成过程转变为生成，评价与修订相结合的闭环过程。",
        None,
    ),
]


def insert_paragraph_after(paragraph: Paragraph, text="", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style and style != "figure":
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def delete_paragraph(paragraph: Paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def add_figure_after(paragraph: Paragraph, image_path: Path, caption: str) -> Paragraph:
    img_p = insert_paragraph_after(paragraph)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.paragraph_format.space_before = Pt(6)
    img_p.paragraph_format.space_after = Pt(2)
    img_p.add_run().add_picture(str(image_path), width=Cm(15.8))
    cap_p = insert_paragraph_after(img_p, caption)
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap_p.runs:
        run.font.size = Pt(9)
    return cap_p


def main():
    make_figure()
    doc = Document(SRC)
    start = end = None
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text == "3 系统架构与方法":
            start = i
        elif start is not None and text.startswith("4 "):
            end = i
            break
    if start is None or end is None:
        raise RuntimeError("Cannot locate chapter 3")

    for idx in range(end - 1, start - 1, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start - 1]
    for text, style in SECTION3:
        if style == "figure":
            anchor = add_figure_after(anchor, FIG, text)
        else:
            anchor = insert_paragraph_after(anchor, text, style)

    if OUT.exists():
        OUT.unlink()
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
