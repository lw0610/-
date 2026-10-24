from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from rewrite_system_architecture_method_academic import FIG_PNG, make_png


ROOT = Path("/Users/zw/研究生/QGIS-agent/Eino")
SRC = ROOT / "AutoCarto_专题图制图智能体论文_系统架构与方法重写版.docx"
OUT = ROOT / "AutoCarto_专题图制图智能体论文_系统架构与方法精简论文版.docx"


SECTION3: list[tuple[str, str | None]] = [
    (
        "本章提出一种面向路线类专题图的知识约束智能制图方法。该方法面向自然语言路线叙事难以直接转化为 GIS 制图操作、通用大模型缺少路线制图语义约束、自动生成地图缺少质量反馈等问题，将路线语义建模、知识图谱约束、多智能体协同、QGIS 工具执行和质量反馈机制组织为统一框架。与直接让大语言模型生成 PyQGIS 脚本不同，本文方法在工具执行前显式构造可检查的制图状态，使路线节点、历史事件、空间位置、表达方法、工具输入和评价结果能够在同一任务链中传递与修订。",
        None,
    ),
    ("图 3-1 知识约束的路线类专题图智能制图方法框架", "figure"),
    ("3.1 总体框架", "Heading 2"),
    (
        "路线类专题图智能制图任务可以定义为从自然语言路线叙事到可验证地图成果的转换过程。给定用户制图指令、历史文献或业务文本，系统需要识别其中的行动主体、时间线索、地点实体、路线节点、历史事件、节点顺序和成果表达要求，并结合已有空间数据与 QGIS 工程状态，生成包含路线线状符号、节点标注、事件说明、图例、比例尺、指北针和输出布局的专题地图。该任务的输入包括文本需求、空间数据、制图知识、模型配置和工具状态；输出包括结构化制图需求、路线制图知识图谱、工具调用序列、QGIS 图层与布局对象以及最终导出的地图成果。",
        None,
    ),
    (
        "AutoCarto AI 采用用户交互层、Go 后端调度层、多智能体层和 QGIS 执行层组成的四层架构。用户交互层负责接收自然语言制图需求、路线文本和制图参数，并展示任务状态、工具日志和地图成果。Go 后端调度层负责维护会话上下文、任务队列、模型配置、工具网关和运行日志，将前端输入转化为可调度的结构化请求。多智能体层负责完成路线意图解析、数据可用性推理、路线关系组织、专题表达设计、版式编排和质量评价。QGIS 执行层以 QGIS/PyQGIS 为本地制图环境，将工具调用转化为图层加载、几何处理、符号化、标注配置、布局生成和成果导出等操作。",
        None,
    ),
    (
        "该总体框架的核心在于区分语言推理、任务调度和 GIS 执行的边界。大语言模型负责理解制图意图、组织路线语义并选择合适工具；Go 后端负责运行时状态维护、工具调用转发和异常记录；QGIS 负责可验证的空间数据处理和地图渲染。通过这种分层设计，系统既避免了模型直接生成长脚本带来的不可控问题，也避免了 GIS 工具孤立执行而缺少语义约束的问题。系统运行过程中形成语义状态、知识状态、执行状态和评价状态四类中间状态，前一状态作为后一模块的输入，并在质量评价后反馈到路线表达和版式编排环节。",
        None,
    ),
    ("3.2 路线制图语义建模与知识组织", "Heading 2"),
    (
        "路线类专题图的核心信息通常以叙事方式出现，例如某一行动主体在特定时间从起点出发，经过若干地点，并在途中发生战斗、转移、驻留或会师等事件。此类文本不能直接作为 GIS 绘图参数使用，因为地点实体、路线顺序、事件归属和表达意图往往隐含在句子结构中。为此，本文将路线制图信息抽取定义为本体约束下的语义结构化过程，即将原始文本解析为可供知识图谱、智能体规划和工具执行共享的制图需求。",
        None,
    ),
    (
        "抽取模式包括行动主体、时间信息、地点实体、路线节点、历史事件、路线顺序、节点重要性、事件说明和成果表达要求等字段。地点实体记录名称、别名、地理线索和地理编码需求；路线节点记录节点序号、节点类型、关联地点和标注等级；历史事件记录事件名称、发生时间、关联节点和说明文本；表达要求记录路线线型、节点符号、说明框、图例和导出格式等需求。模型输出需要满足预定义字段集合和类型约束，并经过字段完整性、节点顺序一致性和事件地点匹配检查。这样，文本中的不确定叙事被转化为可检查的语义状态，从而减少后续工具调用中的字段漂移和参数缺失。",
        None,
    ),
    (
        "在语义抽取基础上，系统构建面向路线类专题图的制图本体与知识图谱。该知识图谱不是通用地理知识库，而是围绕路线图自动制图任务组织的领域知识结构，用于显式描述制图任务、路线节点、路线段、历史事件、空间位置、行动主体、表达方法、地图要素、质量约束和输出成果之间的关系。典型关系包括：制图任务包含路线节点，路线节点按照顺序连接形成路线段，历史事件发生于路线节点，行动主体参与历史事件，路线段采用特定线状表达方法，重要事件需要说明框，地图成果包含标题、图例、比例尺和指北针，质量约束检查路线连贯性、节点完整性、事件匹配和版式可读性。通过这些关系，系统能够在工具执行前明确哪些节点需要连接、哪些事件需要标注、哪些符号需要进入图例以及哪些质量问题需要检查。",
        None,
    ),
    ("3.3 多智能体协同与 QGIS 工具执行", "Heading 2"),
    (
        "路线类专题图制图链条较长，涉及文本理解、数据检查、路线组织、符号设计、版式编排和质量评价等多类知识。若由单一模型一次性完成，容易出现职责混淆、执行顺序错误和质量问题难以定位等情况。本文依据路线类专题图制图过程设置中央协调智能体、路线意图解析智能体、数据推理智能体、路线表达智能体、版式编排智能体和质量评价智能体，使不同知识类型由相应角色处理。",
        None,
    ),
    (
        "中央协调智能体维护全局目标、任务状态和子任务顺序；路线意图解析智能体识别制图目标、行动主体、路线范围和成果形式；数据推理智能体检查路线节点、坐标字段、图层状态和 QGIS 工程环境；路线表达智能体依据图谱关系生成主路线、分支路线、多起点汇合路线和事件节点表达方案；版式编排智能体组织标题、图例、比例尺、指北针、节点标注和说明框；质量评价智能体根据知识图谱、QGIS 状态和地图图像输出问题定位与修订建议。各智能体之间传递的是结构化制图需求、图谱关系、工具执行结果和评价状态，而不是仅依赖自由文本对话。该机制使系统能够把复杂制图任务拆分为可追踪的局部决策，并为实验中的错误归因和消融分析提供基础。",
        None,
    ),
    (
        "QGIS 制图执行采用工具契约而非自由脚本生成的方式组织。每个工具均限定输入参数、输出结果、适用条件和失败反馈，并由 Go 后端统一调度。路线类专题图工具链包括数据接入与检查、路线关系构建、路线表达、地图要素配置和成果导出五类能力。数据接入与检查工具用于加载路线节点、路线段、行政边界和底图数据，并检查坐标系统、字段名称和几何类型；路线关系构建工具依据节点顺序、分支关系和多起点汇合关系生成路线段；路线表达工具生成平滑路线、曲线箭头、主路线与分支路线符号，并处理路线叠压和层级关系；地图要素配置工具生成节点标注、事件说明框、图例、标题、比例尺和指北针；成果导出工具输出 PDF、PNG 或布局工程文件。工具执行过程中产生的图层状态、日志、布局对象和导出成果会被写入执行状态，供质量评价模块进一步检查。",
        None,
    ),
    ("3.4 地图质量评价与反馈优化", "Heading 2"),
    (
        "路线类专题图的质量不仅取决于工具是否成功执行，还取决于地图是否准确表达路线过程、节点是否完整、事件注记是否对应、路线层次是否清晰以及版式是否便于阅读。为避免系统停留在“能生成地图但难以判断质量”的阶段，本文将用户原始需求、路线知识图谱、QGIS 图层与布局状态以及导出的地图图像共同作为质量评价输入，围绕路线连贯性、节点完整性、事件标注准确性、符号层次清晰度、图例一致性、标注可读性和版式完整性等维度开展检查。",
        None,
    ),
    (
        "评价过程采用规则检查与多模态模型判断相结合的方式。对于节点数量、图例项、地图要素完整性和路线顺序等可结构化检查内容，系统优先依据知识图谱和 QGIS 状态进行规则判断；对于视觉层次、标注遮挡、说明框位置和整体可读性等难以完全规则化的问题，系统将地图图像和评价提示输入多模态模型，生成问题描述和修改建议。评价结果进一步被转化为修订指令，反馈给路线表达智能体或版式编排智能体。例如，关键节点缺失时调用节点标注工具补充标注；主路线与分支路线层次不清时调整线宽、线型或颜色；说明框遮挡路线时重新计算说明框位置。由此，系统从一次性地图生成扩展为“生成—评价—修订”的闭环制图流程。",
        None,
    ),
    (
        "该反馈机制也是本文方法区别于普通 GIS 自动化流程的关键环节。普通自动化流程通常以地图导出作为任务结束点，而本文将导出成果重新纳入评价状态，并利用图谱关系、工程状态和地图图像共同定位问题。这样既能保留 QGIS 在空间数据处理和制图表达上的确定性，又能利用多模态模型对视觉可读性和版式合理性进行补充判断。后续实验可从任务完成率、工具调用成功率、问题修复率、人工干预次数和成果完整性等方面验证该闭环机制的有效性。",
        None,
    ),
]


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def delete_paragraph(paragraph: Paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def add_figure_after(paragraph: Paragraph, image_path: Path, caption: str) -> Paragraph:
    img_p = insert_paragraph_after(paragraph)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.paragraph_format.space_before = Pt(6)
    img_p.paragraph_format.space_after = Pt(2)
    img_p.add_run().add_picture(str(image_path), width=Cm(15.8))
    cap_p = insert_paragraph_after(img_p, caption)
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.paragraph_format.space_after = Pt(8)
    for run in cap_p.runs:
        run.font.size = Pt(9)
    return cap_p


def locate_section(doc: Document):
    start = end = None
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text == "3 系统架构与方法":
            start = i
        elif start is not None and text.startswith("4 "):
            end = i
            break
    if start is None or end is None:
        raise RuntimeError("Cannot locate section 3")
    return start, end


def main():
    make_png()
    doc = Document(SRC)
    start, end = locate_section(doc)
    for idx in range(end - 1, start, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start]
    for text, kind in SECTION3:
        if kind == "figure":
            anchor = add_figure_after(anchor, FIG_PNG, text)
        else:
            anchor = insert_paragraph_after(anchor, text, style=kind)

    if OUT.exists():
        OUT.unlink()
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
